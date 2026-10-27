import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from uuid import uuid4
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app
from app.testcase_export import set_text


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app(Path(self.temp.name) / 'export.sqlite3')
        self.client = TestClient(self.app)
        self.client.__enter__()
        self.pid = self.client.post('/api/projects', json={'name': '导出/验收项目'}).json()['id']
        root = f'/api/projects/{self.pid}'
        rid = self.client.post(root + '/requirements', json={'title': '登录', 'content': '账号登录规则'}).json()['id']
        self.requirement = f'{root}/requirements/{rid}'
        sheet = self.client.get(self.requirement + '/features').json()
        sheet = self.client.put(self.requirement + '/features', json={
            'expected_revision': 0, 'input_hash': sheet['current_hash'], 'confirmed': True,
            'items': [{'module': '账号', 'feature': '登录', 'description': '账号登录'}],
        }).json()
        self.client.put('/api/settings/generation', json={'mock_mode': True})
        generated = self.client.post(self.requirement + '/case-generations', json={
            'request_id': str(uuid4()), 'feature_id': sheet['items'][0]['id'],
            'expected_revision': sheet['revision'], 'input_hash': sheet['current_hash'], 'strategy': 'full',
        })
        self.assertEqual(generated.status_code, 201, generated.text)
        generation = generated.json()
        self.review = self.requirement + '/case-generations/' + generation['id'] + '/reviews/'
        self.ids = [item['id'] for item in generation['cases']]
        current = self.client.get(self.review.rstrip('/')).json()['items'][0]['current']
        self.fields = {**current, 'title': '=HYPERLINK("https://example.invalid", "测试")',
                       'precondition': '保留控制符\x01与换行\n准备',
                       'steps': ['输入账号', '```\n<script>alert(1)</script>\n```'],
                       'expected_result': '+正常登录\n显示结果'}
        self.assertEqual(self.client.put(self.review + self.ids[0], json={
            'action': 'edit', 'expected_revision': 0, 'fields': self.fields}).status_code, 200)
        for index in (0, 1):
            self.assertEqual(self.client.put(self.review + self.ids[index], json={
                'action': 'accept', 'expected_revision': 1 if index == 0 else 0}).status_code, 200)
        self.endpoint = root + '/testcases/export'

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.temp.cleanup()

    def export(self, format='xlsx', ids=None, endpoint=None):
        return self.client.post(endpoint or self.endpoint, json={
            'format': format, 'case_ids': self.ids[:2] if ids is None else ids})

    def test_excel_edited_content_text_safety_layout_and_order(self):
        response = self.export(ids=list(reversed(self.ids[:2])))
        self.assertEqual(response.status_code, 200, response.text[:300])
        self.assertEqual(response.headers['cache-control'], 'no-store')
        filename = unquote(response.headers['content-disposition'].split("filename*=UTF-8''")[1])
        self.assertTrue(filename.startswith('导出_验收项目_已采纳用例_'))
        self.assertTrue(filename.endswith('.xlsx'))
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 6)
        self.assertEqual(sheet.max_column, 15)
        self.assertEqual(sheet['A5'].value, self.ids[1])
        self.assertEqual(sheet['A6'].value, self.ids[0])
        self.assertEqual(sheet['E6'].value, self.fields['title'])
        self.assertEqual(sheet['E6'].data_type, 's')
        self.assertEqual(sheet['I6'].value, '保留控制符\\u0001与换行\n准备')
        self.assertEqual(sheet['J6'].value, '1. 输入账号\n2. ```\n<script>alert(1)</script>\n```')
        self.assertEqual(sheet['K6'].value, self.fields['expected_result'])
        self.assertEqual(sheet['K6'].data_type, 's')
        self.assertEqual(sheet['L6'].value, 'Mock 示例')
        self.assertEqual(sheet['M6'].value, '是')
        self.assertEqual(sheet.freeze_panes, 'A5')
        self.assertEqual(sheet.auto_filter.ref, 'A4:O6')
        self.assertTrue(sheet['J6'].alignment.wrap_text)
        with ZipFile(BytesIO(response.content)) as archive:
            self.assertNotIn(b'<f>', archive.read('xl/worksheets/sheet1.xml'))

    def test_markdown_filtered_current_content_and_safe_fences(self):
        response = self.export('markdown', [self.ids[0]])
        self.assertEqual(response.status_code, 200)
        self.assertIn('text/markdown', response.headers['content-type'])
        text = response.content.decode('utf-8')
        self.assertIn('- 用例数：1', text)
        self.assertIn(self.ids[0], text)
        self.assertNotIn(self.ids[1], text)
        self.assertIn('1. 输入账号', text)
        self.assertIn('````text\n1. 输入账号\n2. ```', text)
        self.assertIn('人工编辑：是', text)
        self.assertIn('Mock 示例', text)
        self.assertIn('\\u0001', text)
        # HTML in user content is in a fenced code block; metadata headings are escaped.
        self.assertIn('<script>alert(1)</script>', text)

    def test_outdated_accepted_cases_are_marked_in_both_formats(self):
        self.client.put(self.requirement, json={'title': '更新需求', 'content': '规则变更'})
        self.assertIn('依据已变化', self.export('markdown').text)
        sheet = load_workbook(BytesIO(self.export().content)).active
        self.assertEqual(sheet['N5'].value, '依据已变化')
        self.assertEqual(sheet['B5'].value, '登录')

    def test_pending_rejected_missing_and_cross_project_ids_fail_atomically(self):
        for format in ('xlsx', 'markdown'):
            self.assertEqual(self.export(format, [self.ids[0], self.ids[2]]).status_code, 409)
        self.client.put(self.review + self.ids[2], json={
            'action': 'reject', 'expected_revision': 0, 'rejection_reason': '不采纳'})
        self.assertEqual(self.export(ids=[self.ids[2]]).status_code, 409)
        self.assertEqual(self.export(ids=[str(uuid4())]).status_code, 409)
        other = self.client.post('/api/projects', json={'name': '其他项目'}).json()['id']
        self.assertEqual(self.export(endpoint=f'/api/projects/{other}/testcases/export').status_code, 409)
        self.assertEqual(self.export(endpoint='/api/projects/999999/testcases/export').status_code, 404)
        self.client.delete(self.requirement)
        self.assertEqual(self.export().status_code, 409)

    def test_invalid_export_requests_are_rejected(self):
        for format, ids in [('csv', self.ids[:1]), ('xlsx', []), ('xlsx', [self.ids[0]] * 2),
                            ('xlsx', ['invalid']), ('xlsx', [str(uuid4()) for _ in range(1001)])]:
            self.assertEqual(self.export(format, ids).status_code, 422)

    def test_export_does_not_change_reviews_or_call_model(self):
        before = self.client.get(self.review.rstrip('/')).json()
        # Turning off Mock without configuring a model proves export is independent of generation.
        self.client.put('/api/settings/generation', json={'mock_mode': False})
        self.assertEqual(self.export().status_code, 200)
        self.assertEqual(self.export('markdown').status_code, 200)
        self.assertEqual(self.client.get(self.review.rstrip('/')).json(), before)

    def test_excel_long_unicode_rejected_instead_of_truncated(self):
        from fastapi import HTTPException
        from openpyxl import Workbook
        cell = Workbook().active['A1']
        with self.assertRaises(HTTPException) as raised:
            set_text(cell, '😀' * 16384)
        self.assertEqual(raised.exception.status_code, 400)


if __name__ == '__main__':
    unittest.main()
