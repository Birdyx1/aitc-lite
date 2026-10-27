"""导出只读的已采纳用例快照，不调用模型、不修改业务数据。"""

import html
import math
import re
from datetime import datetime, timedelta, timezone
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from .deps import get_session
from .projects import require_project
from .reviews import list_accepted
from .schemas import AcceptedCaseOutput, CaseExportInput

router = APIRouter(prefix="/api/projects/{project_id}/testcases", tags=["用例导出"])
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TYPE_NAMES = {"functional": "功能", "boundary": "边界", "exception": "异常"}
HEADERS = ["用例ID", "需求", "模块", "功能点", "用例标题", "优先级", "类型", "冒烟", "前置条件", "操作步骤", "预期结果", "来源", "人工编辑", "依据状态", "生成记录ID"]
WIDTHS = [39, 24, 18, 26, 38, 10, 10, 10, 34, 52, 52, 14, 14, 20, 39]


def safe_text(value) -> str:
    # 将 XML/UTF-8 不可表示字符变成可见转义，不静默删除。
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]",
                  lambda match: f"\\u{ord(match.group()):04x}", str(value))


def steps_text(item: AcceptedCaseOutput) -> str:
    return "\n".join(f"{index}. {step}" for index, step in enumerate(item.steps, 1))


def row_values(item: AcceptedCaseOutput) -> list[str]:
    return [str(item.id), item.requirement_title, item.module, item.feature_name, item.title,
            item.priority, TYPE_NAMES[item.case_type], "是" if item.is_smoke else "否", item.precondition,
            steps_text(item), item.expected_result, "Mock 示例" if item.source == "mock" else "AI 生成",
            "是" if item.was_edited else "否", "依据已变化" if item.outdated else "与当前输入一致", str(item.generation_id)]


def set_text(cell, value):
    text = safe_text(value)
    if len(text.encode("utf-16-le")) // 2 > 32767:
        raise HTTPException(status_code=400, detail="个别用例超过 Excel 单元格文本上限，请缩短内容或改用 Markdown 导出。")
    cell.value = text
    cell.data_type = "s"  # 包括 =、+、-、@ 开头的用户内容，始终作为文本写入。


def export_xlsx(project_name: str, cases: list[AcceptedCaseOutput], timestamp: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "已采纳用例"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:O1")
    set_text(sheet["A1"], f"{project_name} · 已采纳测试用例")
    sheet["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color="23344D")
    sheet.row_dimensions[1].height = 36
    sheet.merge_cells("A2:O2")
    set_text(sheet["A2"], f"导出时间：{timestamp}  |  用例数：{len(cases)}  |  范围：页面当前筛选结果，均为已采纳用例")
    sheet["A2"].font = Font(name="Microsoft YaHei", size=11, color="68778D")
    sheet.row_dimensions[2].height = 26
    sheet.merge_cells("A3:O3")
    set_text(sheet["A3"], "保留生成来源及依据状态；超长单元格可在公式栏查看全文。不可表示字符以可见 Unicode 转义显示。")
    sheet["A3"].font = Font(name="Microsoft YaHei", size=10, color="8793A4")
    for index, header in enumerate(HEADERS, 1):
        cell = sheet.cell(4, index)
        set_text(cell, header)
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="354B74")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS[index - 1]
    sheet.row_dimensions[4].height = 28
    for row_index, item in enumerate(cases, 5):
        values = row_values(item)
        lines = 1
        for index, value in enumerate(values, 1):
            cell = sheet.cell(row_index, index)
            set_text(cell, value)
            cell.font = Font(name="Microsoft YaHei", size=11, color="344257")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="F4F7FC" if row_index % 2 else "FFFFFF")
            if (index == 12 and item.source == "mock") or (index == 14 and item.outdated):
                cell.font = Font(name="Microsoft YaHei", size=11, color="946515")
                cell.fill = PatternFill("solid", fgColor="FFF4D9")
            # 中文通常占两字符宽度，估算换行高度，长内容保留不截断。
            estimated = sum(max(1, math.ceil(len(line) * 1.6 / WIDTHS[index - 1])) for line in safe_text(value).split("\n"))
            lines = max(lines, estimated)
        sheet.row_dimensions[row_index].height = min(390, max(72, lines * 16 + 12))
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:O{sheet.max_row}"
    sheet.print_title_rows = "1:4"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def inline(value) -> str:
    text = html.escape(safe_text(value), quote=False)
    text = re.sub(r"([\\`*_{}\[\]()#+.!|>\-])", r"\\\1", text)
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")


def text_block(value: str) -> str:
    text = safe_text(value)
    longest = max((len(match) for match in re.findall(r"`+", text)), default=0)
    fence = "`" * max(3, longest + 1)
    return f"{fence}text\n{text}\n{fence}\n"


def export_markdown(project_name: str, cases: list[AcceptedCaseOutput], timestamp: str) -> bytes:
    parts = ["# 已采纳测试用例\n", f"- 项目：{inline(project_name)}\n- 导出时间：{timestamp}\n- 用例数：{len(cases)}\n",
             "范围：页面当前筛选结果，均为已采纳用例。保留来源、人工编辑和依据状态；正文按原文文本保存。\n"]
    for index, item in enumerate(cases, 1):
        parts.extend([
            f"## {index}. {inline(item.title)}\n",
            f"- 用例 ID：`{item.id}`\n- 需求：{inline(item.requirement_title)}\n- 模块：{inline(item.module)}\n- 功能点：{inline(item.feature_name)}\n"
            f"- 优先级：{item.priority}\n- 类型：{TYPE_NAMES[item.case_type]}\n- 冒烟：{'是' if item.is_smoke else '否'}\n"
            f"- 来源：{'Mock 示例' if item.source == 'mock' else 'AI 生成'}\n- 人工编辑：{'是' if item.was_edited else '否'}\n"
            f"- 依据状态：{'依据已变化' if item.outdated else '与当前输入一致'}\n- 生成记录 ID：`{item.generation_id}`\n",
            "### 前置条件\n", text_block(item.precondition or "无"),
            "### 操作步骤\n", text_block(steps_text(item)),
            "### 预期结果\n", text_block(item.expected_result),
        ])
    return "\n".join(parts).encode("utf-8")


@router.post("/export")
def export_cases(project_id: int, payload: CaseExportInput, session: Session = Depends(get_session)):
    project = require_project(session, project_id)
    allowed = {str(item.id): item for item in list_accepted(project_id, session)}
    ids = [str(case_id) for case_id in payload.case_ids]
    if any(case_id not in allowed for case_id in ids):
        raise HTTPException(status_code=409, detail="所选用例已删除或不在当前项目已采纳范围，请刷新列表后重新导出。")
    cases = [allowed[case_id] for case_id in ids]
    now = datetime.now(timezone(timedelta(hours=8)))
    timestamp = now.strftime("%Y-%m-%d %H:%M:%S UTC+08:00")
    xlsx = payload.format == "xlsx"
    data = export_xlsx(project.name, cases, timestamp) if xlsx else export_markdown(project.name, cases, timestamp)
    name = re.sub(r'[\x00-\x1f\x7f\\/:*?"<>|\ud800-\udfff]', "_", project.name).strip(" .")[:80] or "AITC"
    suffix = "xlsx" if xlsx else "md"
    filename = f"{name}_已采纳用例_{now:%Y%m%d_%H%M%S}.{suffix}"
    return Response(data, media_type=XLSX_TYPE if xlsx else "text/markdown", headers={
        "Content-Disposition": f"attachment; filename=\"aitc-testcases.{suffix}\"; filename*=UTF-8''{quote(filename, safe='')}",
        "Cache-Control": "no-store",
    })
